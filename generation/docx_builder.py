from docx import Document
import re

def _extract_contact_fallback(cv_text: str) -> dict:
    """
    Fallback: pulls email/phone from raw CV text via regex if not
    explicitly provided.
    """
    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', cv_text)
    phone_match = re.search(r'(\+?\d[\d\s\-\(\)]{7,}\d)', cv_text)
    return {
        "email": email_match.group(0) if email_match else "",
        "phone": phone_match.group(0) if phone_match else ""
    }

def build_tailored_resume(candidate_name: str, tailored_summary: str,
                            reordered_skills: list, tailored_bullets: list,
                            experience: list, education: list, certifications: list,
                            cv_text: str, output_path: str,
                            contact_email: str = None, contact_phone: str = None) -> str:
    """
    Assembles a tailored resume into a real .docx file, using only
    verified real content — no new content is generated here.

    Args:
        candidate_name: the candidate's name (must be provided)
        tailored_summary: from tailor_summary()
        reordered_skills: from tailor_skills()
        tailored_bullets: from tailor_experience()
        experience: from analyze_cv()['experience'] — real, untouched
        education: from analyze_cv()['education'] — real, untouched
        certifications: from analyze_cv()['certifications'] — real, untouched
        cv_text: raw CV text, used only as a fallback to find contact info
        output_path: where to save the .docx file
        contact_email, contact_phone: optional — falls back to regex extraction

    Returns:
        The output_path, once the file is saved
    """
    if not contact_email or not contact_phone:
        fallback = _extract_contact_fallback(cv_text)
        contact_email = contact_email or fallback['email']
        contact_phone = contact_phone or fallback['phone']

    doc = Document()

    doc.add_heading(candidate_name, level=1)
    doc.add_paragraph(f"{contact_email}  |  {contact_phone}")

    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(tailored_summary)

    doc.add_heading("Skills", level=2)
    doc.add_paragraph(", ".join(reordered_skills))

    doc.add_heading("Key Highlights", level=2)
    for bullet in tailored_bullets:
        doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Professional Experience", level=2)
    for job in experience:
        line = f"{job.get('title', '')} — {job.get('company', '')}"
        if job.get('duration'):
            line += f" ({job['duration']})"
        doc.add_paragraph(line)

    doc.add_heading("Education", level=2)
    for edu in education:
        line = f"{edu.get('degree', '')} — {edu.get('institution', '')}"
        if edu.get('year'):
            line += f" ({edu['year']})"
        doc.add_paragraph(line)

    if certifications:
        doc.add_heading("Certifications", level=2)
        doc.add_paragraph(", ".join(certifications))

    doc.save(output_path)
    return output_path